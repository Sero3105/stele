"""Generate the benchmark/test dataset with known ground truth.

Every PDF/image is rendered from text we control, so OCR output can be scored
(CER/WER) against `<name>.gt.txt`. Scans are simulated by rasterising the
digital PDF and degrading it (resolution, noise, blur, rotation, JPEG).

Run: .venv/bin/python benchmarks/make_dataset.py [out_dir]
"""

from __future__ import annotations

import io
import random
import sys
from pathlib import Path

import numpy as np
import pypdfium2 as pdfium
from PIL import Image, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate, Frame, ListFlowable, PageTemplate, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent / "data")
STYLES = getSampleStyleSheet()

IT_TITLE = "Relazione tecnica sul consumo energetico"
IT_SECTIONS = [
    ("Introduzione", [
        "Questa relazione descrive il consumo energetico degli edifici comunali nel corso dell'anno. "
        "L'analisi è stata condotta su dodici strutture, tra cui scuole, uffici e impianti sportivi.",
        "I dati sono stati raccolti mensilmente dai contatori elettrici e termici. Ogni valore è stato "
        "verificato confrontando le letture con le fatture dei fornitori.",
    ]),
    ("Metodologia", [
        "Per ciascun edificio abbiamo calcolato l'intensità energetica, cioè il rapporto tra energia "
        "consumata e superficie utile. Questo indicatore permette di confrontare strutture di dimensioni diverse.",
        "Le misure anomale, dovute a guasti o letture mancanti, sono state sostituite con la media dei mesi adiacenti.",
    ]),
]
IT_TABLE = [
    ["Edificio", "Superficie (m²)", "Consumo (kWh)", "Intensità"],
    ["Scuola Rossi", "2450", "182300", "74,4"],
    ["Municipio", "3120", "251900", "80,7"],
    ["Palestra Nord", "1800", "96400", "53,6"],
    ["Biblioteca", "950", "41200", "43,4"],
]
IT_SECTIONS_2 = [
    ("Risultati", [
        "La palestra e la biblioteca mostrano i valori più bassi grazie ai recenti lavori di isolamento. "
        "Il municipio, invece, presenta l'intensità più elevata a causa degli impianti di climatizzazione obsoleti.",
    ]),
    ("Conclusioni", [
        "Si raccomanda di sostituire le caldaie del municipio entro il prossimo esercizio e di installare "
        "sensori di presenza negli uffici. Il risparmio stimato è pari al diciotto per cento.",
    ]),
]
IT_LIST = [
    "Sostituzione delle caldaie con pompe di calore.",
    "Installazione di pannelli fotovoltaici sulle coperture.",
    "Monitoraggio continuo dei consumi con report trimestrali.",
]

EN_TITLE = "Field Notes on Coastal Erosion Monitoring"
EN_PARAS = [
    "Coastal erosion is measured by comparing shoreline positions over time. Our team surveyed four beaches "
    "between March and October, using fixed markers and satellite positioning.",
    "The northern beach lost an average of 1.8 metres of width, mostly after two storms in late September. "
    "The southern beaches were stable, protected by a natural reef that dissipates wave energy.",
    "Sediment samples were collected at low tide along transects perpendicular to the shore. Grain size "
    "analysis showed coarser sand where wave action was strongest.",
    "Volunteers from local schools helped record observations, which improved coverage on weekends and "
    "holidays when professional staff were not available.",
    "We recommend repeating the survey every six months and adding drone photogrammetry to capture changes "
    "in dune height that ground measurements may miss.",
    "All raw measurements are archived with their timestamps so that future studies can reproduce the analysis "
    "and extend the time series.",
]
EN_TABLE = [
    ["Beach", "Width change (m)", "Storm events"],
    ["North", "-1.8", "2"],
    ["Central", "-0.6", "1"],
    ["South A", "0.1", "0"],
    ["South B", "0.0", "0"],
]

FORMULA_LINES = [
    "Equations used in the model",
    "The kinetic energy of a body is E = 1/2 m v^2, where m is the mass and v the velocity.",
    "The area of a circle is A = π r^2.",
    "The roots of a x^2 + b x + c = 0 are x = (-b ± √(b^2 - 4ac)) / 2a.",
    "The sum of the first n integers is n(n+1)/2.",
]


def chart_png() -> Path:
    """A small bar chart (no text inside, so OCR ground truth is unchanged)."""
    from PIL import ImageDraw
    path = OUT / "_chart.png"
    im = Image.new("RGB", (480, 300), "white")
    d = ImageDraw.Draw(im)
    for i, h in enumerate([182, 252, 96, 41]):
        d.rectangle([40 + i * 110, 280 - h, 120 + i * 110, 280], fill=(40, 90 + i * 40, 160))
    d.line([30, 280, 470, 280], fill="black", width=2)
    im.save(path)
    return path


CHART_TITLE = "Consumo annuo per edificio (MWh)"
CHART_DATA = [("Scuola", 182), ("Municipio", 252), ("Palestra", 96), ("Biblioteca", 41)]


def labelled_chart_png() -> Path:
    """Bar chart with title, category labels and the value printed on each bar."""
    from PIL import ImageDraw, ImageFont
    f = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 22)
    fs = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 18)
    im = Image.new("RGB", (900, 560), "white")
    d = ImageDraw.Draw(im)
    d.text((250, 15), CHART_TITLE, fill="black", font=f)
    for i, (name, v) in enumerate(CHART_DATA):
        x, h = 100 + i * 190, v * 1.6
        d.rectangle([x, 480 - h, x + 110, 480], fill=(40, 100 + i * 30, 170))
        d.text((x + 35, 480 - h - 28), str(v), fill="black", font=fs)
        d.text((x + 10, 490), name, fill="black", font=fs)
    d.line([80, 480, 860, 480], fill="black", width=2)
    path = OUT / "_labelled_chart.png"
    im.save(path)
    return path


def make_chart_pdf() -> None:
    from reportlab.platypus import Image as RLImage
    story = [Paragraph("Andamento dei consumi", STYLES["Heading2"]),
             Paragraph("Il grafico seguente riassume i consumi misurati.", STYLES["BodyText"]),
             RLImage(str(labelled_chart_png()), width=14 * cm, height=8.7 * cm),
             Paragraph("Figura 1: consumo annuo per edificio.", STYLES["BodyText"])]
    SimpleDocTemplate(str(OUT / "digital_chart_it.pdf"), pagesize=A4).build(story)
    (OUT / "_labelled_chart.png").unlink()


def table_flowable(rows: list[list[str]]) -> Table:
    t = Table(rows, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ]))
    return t


def write_gt(name: str, chunks: list[str]) -> None:
    (OUT / f"{name}.gt.txt").write_text("\n".join(chunks) + "\n", encoding="utf-8")


def table_gt(rows: list[list[str]]) -> list[str]:
    return [" ".join(r) for r in rows]


def make_italian_pdf() -> None:
    story, gt = [Paragraph(IT_TITLE, STYLES["Title"])], [IT_TITLE]
    for head, paras in IT_SECTIONS:
        story.append(Paragraph(head, STYLES["Heading2"])); gt.append(head)
        for p in paras:
            story.append(Paragraph(p, STYLES["BodyText"])); gt.append(p)
    story += [Spacer(1, 0.4 * cm), table_flowable(IT_TABLE)]
    gt += table_gt(IT_TABLE)
    from reportlab.platypus import PageBreak
    story.append(PageBreak())
    for head, paras in IT_SECTIONS_2:
        story.append(Paragraph(head, STYLES["Heading2"])); gt.append(head)
        for p in paras:
            story.append(Paragraph(p, STYLES["BodyText"])); gt.append(p)
    story.append(ListFlowable([Paragraph(i, STYLES["BodyText"]) for i in IT_LIST], bulletType="bullet"))
    gt += IT_LIST
    from reportlab.platypus import Image as RLImage
    story += [Spacer(1, 0.4 * cm), RLImage(str(chart_png()), width=8 * cm, height=5 * cm)]
    SimpleDocTemplate(str(OUT / "digital_it.pdf"), pagesize=A4, title=IT_TITLE).build(story)
    write_gt("digital_it", gt)


def make_english_two_column_pdf() -> None:
    doc = BaseDocTemplate(str(OUT / "digital_en_twocol.pdf"), pagesize=A4, title=EN_TITLE)
    w, h = A4
    m, gap = 2 * cm, 0.8 * cm
    colw = (w - 2 * m - gap) / 2
    title_h = 3 * cm
    frames = [
        Frame(m, h - m - title_h, w - 2 * m, title_h, id="title"),
        Frame(m, m, colw, h - 2 * m - title_h, id="left"),
        Frame(m + colw + gap, m, colw, h - 2 * m - title_h, id="right"),
    ]
    doc.addPageTemplates([PageTemplate(frames=frames)])
    from reportlab.platypus import FrameBreak
    story = [Paragraph(EN_TITLE, STYLES["Title"]), FrameBreak()]
    for i, p in enumerate(EN_PARAS):
        story.append(Paragraph(p, STYLES["BodyText"]))
        if i == 2:
            story.append(FrameBreak())
    story += [Spacer(1, 0.3 * cm), table_flowable(EN_TABLE)]
    doc.build(story)
    write_gt("digital_en_twocol", [EN_TITLE, *EN_PARAS, *table_gt(EN_TABLE)])


def make_formula_pdf() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    # Helvetica lacks π/√/±: use a system TrueType font that has them.
    font = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
    pdfmetrics.registerFont(TTFont("ArialUni", font))
    style = STYLES["BodyText"].clone("uni", fontName="ArialUni", fontSize=12, leading=18)
    story = [Paragraph(FORMULA_LINES[0], STYLES["Heading2"])]
    story += [Paragraph(line, style) for line in FORMULA_LINES[1:]]
    SimpleDocTemplate(str(OUT / "digital_formula_en.pdf"), pagesize=A4).build(story)
    write_gt("digital_formula_en", FORMULA_LINES)


def rasterise(pdf: Path, dpi: int) -> list[Image.Image]:
    doc = pdfium.PdfDocument(pdf)
    try:
        return [doc[i].render(scale=dpi / 72).to_pil().convert("L") for i in range(len(doc))]
    finally:
        doc.close()


def degrade(img: Image.Image, *, noise: float, blur: float, rotate: float, jpeg: int, seed: int) -> Image.Image:
    rng = np.random.default_rng(seed)
    if rotate:
        img = img.rotate(rotate, resample=Image.BICUBIC, expand=False, fillcolor=255)
    if blur:
        img = img.filter(ImageFilter.GaussianBlur(blur))
    arr = np.asarray(img, dtype=np.float32)
    arr = np.clip(arr + rng.normal(0, noise, arr.shape), 0, 255).astype(np.uint8)
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, "JPEG", quality=jpeg)
    return Image.open(io.BytesIO(buf.getvalue())).convert("L")


def save_image_pdf(images: list[Image.Image], path: Path, dpi: int) -> None:
    images[0].save(path, "PDF", save_all=True, append_images=images[1:], resolution=dpi)


SCAN_PROFILES = {
    "clean300": dict(dpi=300, noise=4, blur=0, rotate=0, jpeg=90),
    "noisy150": dict(dpi=150, noise=14, blur=0.6, rotate=1.2, jpeg=45),
    "poor100": dict(dpi=100, noise=22, blur=1.0, rotate=-2.0, jpeg=30),
}


def make_scans() -> None:
    for src in ("digital_it", "digital_en_twocol"):
        gt = (OUT / f"{src}.gt.txt").read_text(encoding="utf-8")
        for name, p in SCAN_PROFILES.items():
            pages = [
                degrade(img, noise=p["noise"], blur=p["blur"], rotate=p["rotate"], jpeg=p["jpeg"], seed=i)
                for i, img in enumerate(rasterise(OUT / f"{src}.pdf", p["dpi"]))
            ]
            out = f"scan_{name}_{src.removeprefix('digital_')}"
            save_image_pdf(pages, OUT / f"{out}.pdf", p["dpi"])
            (OUT / f"{out}.gt.txt").write_text(gt, encoding="utf-8")
    # Standalone images.
    page = rasterise(OUT / "digital_it.pdf", 200)[0]
    page.save(OUT / "image_it_table.png")
    (OUT / "image_it_table.gt.txt").write_text(
        "\n".join([IT_TITLE, *[x for h, ps in IT_SECTIONS for x in (h, *ps)], *table_gt(IT_TABLE)]) + "\n",
        encoding="utf-8",
    )
    en = degrade(rasterise(OUT / "digital_en_twocol.pdf", 200)[0], noise=10, blur=0.4, rotate=0.8, jpeg=60, seed=7)
    en.save(OUT / "image_en_photo.jpg", quality=70)
    (OUT / "image_en_photo.gt.txt").write_text((OUT / "digital_en_twocol.gt.txt").read_text(), encoding="utf-8")


def make_mixed_pdf() -> None:
    """Pages: 1-2 digital (IT), 3 scanned (EN two-column), 4 digital (formulas)."""
    out = pdfium.PdfDocument.new()
    for name in ("digital_it.pdf", "scan_noisy150_en_twocol.pdf", "digital_formula_en.pdf"):
        src = pdfium.PdfDocument(OUT / name)
        out.import_pages(src)
        src.close()
    out.save(OUT / "mixed_digital_scan.pdf")
    out.close()
    gt = [(OUT / f"{n}.gt.txt").read_text(encoding="utf-8") for n in ("digital_it", "digital_en_twocol", "digital_formula_en")]
    (OUT / "mixed_digital_scan.gt.txt").write_text("".join(gt), encoding="utf-8")


LONG_SENTENCES = [
    "Il comune ha avviato un programma pluriennale di manutenzione degli edifici pubblici.",
    "Ogni intervento viene pianificato in base allo stato delle strutture e alle risorse disponibili.",
    "I tecnici compilano una scheda per ciascun sopralluogo, indicando i problemi riscontrati.",
    "Le priorità sono stabilite considerando la sicurezza, il risparmio energetico e il comfort degli utenti.",
    "Nel corso dell'estate sono stati sostituiti gli infissi di tre scuole primarie.",
    "La biblioteca centrale ha ricevuto un nuovo impianto di illuminazione a basso consumo.",
    "I lavori sono stati eseguiti senza interrompere i servizi principali per i cittadini.",
    "Il monitoraggio dei consumi continuerà per verificare l'efficacia degli interventi realizzati.",
    "Una parte dei fondi proviene da contributi regionali dedicati alla transizione energetica.",
    "Il prossimo rapporto illustrerà i risultati ottenuti e le attività previste per l'anno successivo.",
]


def make_long_scan(pages: int = 12) -> None:
    """Prose-only scanned report: the common 'long scanned PDF' case (no tables)."""
    rng = random.Random(42)
    story, gt = [], []
    from reportlab.platypus import PageBreak
    for n in range(1, pages + 1):
        head = f"Capitolo {n}"
        story.append(Paragraph(head, STYLES["Heading2"]))
        gt.append(head)
        for _ in range(5):
            para = " ".join(rng.sample(LONG_SENTENCES, 4))
            story.append(Paragraph(para, STYLES["BodyText"]))
            gt.append(para)
        story.append(PageBreak())
    SimpleDocTemplate(str(OUT / "_long.pdf"), pagesize=A4).build(story)
    p = SCAN_PROFILES["noisy150"]
    imgs = [degrade(img, noise=p["noise"], blur=p["blur"], rotate=p["rotate"], jpeg=p["jpeg"], seed=i)
            for i, img in enumerate(rasterise(OUT / "_long.pdf", p["dpi"]))]
    save_image_pdf(imgs, OUT / "scan_long_it.pdf", p["dpi"])
    (OUT / "_long.pdf").unlink()
    write_gt("scan_long_it", gt)


def make_office() -> None:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.add_heading(IT_TITLE, 0)
    for head, paras in IT_SECTIONS:
        d.add_heading(head, 1)
        for p in paras:
            d.add_paragraph(p)
    t = d.add_table(rows=len(IT_TABLE), cols=len(IT_TABLE[0]))
    t.style = "Table Grid"
    for r, row in enumerate(IT_TABLE):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_picture(str(chart_png()), width=docx.shared.Cm(8))
    cap = d.add_paragraph("Tabella 1: consumi per edificio")
    cap.runs[0].font.size = Pt(9)
    for item in IT_LIST:
        d.add_paragraph(item, style="List Bullet")
    d.add_paragraph("Primo passo", style="List Number")
    d.add_paragraph("Secondo passo", style="List Number")
    d.save(OUT / "office_report_it.docx")

    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = EN_TITLE
    s.placeholders[1].text = "Survey results, autumn season"
    s.notes_slide.notes_text_frame.text = "Speaker note: thank the volunteers."
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Key findings"
    s.placeholders[1].text = "North beach lost 1.8 m\nSouthern beaches stable\nReef dissipates wave energy"
    s.notes_slide.notes_text_frame.text = "Mention the September storms."
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Measurements"
    shape = s.shapes.add_table(len(EN_TABLE), 3, Inches(0.5), Inches(1.5), Inches(8), Inches(2.5))
    for r, row in enumerate(EN_TABLE):
        for c, v in enumerate(row):
            shape.table.cell(r, c).text = v
    prs.save(OUT / "office_slides_en.pptx")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consumi"
    ws.append(["Riepilogo consumi 2025"])
    ws.merge_cells("A1:D1")
    for row in IT_TABLE:
        ws.append([float(v.replace(",", ".")) if v.replace(",", "").isdigit() and i else v for i, v in enumerate(row)])
    ws.append(["Totale", "=SUM(B3:B6)", "=SUM(C3:C6)", ""])
    ws2 = wb.create_sheet("Spiagge")
    for row in EN_TABLE:
        ws2.append(row)
    wb.save(OUT / "office_data.xlsx")


def make_misc() -> None:
    (OUT / "note.txt").write_text("Promemoria\n\nRicontrollare i contatori della palestra a novembre.\n", encoding="utf-8")
    (OUT / "dati.csv").write_text("edificio,consumo\nScuola Rossi,182300\nMunicipio,251900\n", encoding="utf-8")
    (OUT / "pagina.html").write_text(
        "<html><body><h1>Avviso</h1><p>La biblioteca resterà chiusa <b>lunedì</b>.</p>"
        "<ul><li>Orari ridotti</li><li>Prestiti online</li></ul></body></html>", encoding="utf-8")
    good = (OUT / "digital_it.pdf").read_bytes()
    (OUT / "broken_truncated.pdf").write_bytes(good[: len(good) // 3])
    (OUT / "unsupported.xyz").write_bytes(bytes(random.Random(1).randrange(256) for _ in range(512)))
    (OUT / "empty.pdf").write_bytes(b"")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_italian_pdf()
    make_english_two_column_pdf()
    make_formula_pdf()
    make_scans()
    make_long_scan()
    make_mixed_pdf()
    make_office()
    make_chart_pdf()
    make_misc()
    (OUT / "_chart.png").unlink()
    print(f"dataset in {OUT}:")
    for p in sorted(OUT.iterdir()):
        print(f"  {p.name:40s} {p.stat().st_size / 1024:8.1f} KB")


if __name__ == "__main__":
    main()
