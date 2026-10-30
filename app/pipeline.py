"""Convert one file: analyse -> route pages -> run engines -> quality check/escalate ->
assemble ONE document -> publish it.

Everything engine-specific stays in the work dir (checkpoints, internal report). The user-facing
result is a single folder: <name>.md, assets/, metadata.json — the same shape whatever mix of
engines produced the pages.
"""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np
from PIL import Image

from app import charts, config, markdown, router
from app.analyzer import FileInfo, analyze, render_pages, soffice
from app.checkpoint import Checkpoint, atomic_write
from app.engines.base import PageOutput
from app.engines.docling_engine import DoclingEngine, page_outputs
from app.engines.vision_engine import VisionEngine
from app.engines.vlm_engine import VlmEngine, transfer_headings
from app.office import pptx_notes, xlsx_details
from app.quality import misspelled_ratio, text_quality
from app.resources import EngineResources
from app.tables import restore_tables

log = logging.getLogger(__name__)

CHUNK_PAGES = 20  # Docling pass size: bounds memory on long PDFs and gives progress updates

# Measured in benchmarks/REPORT.md (model = resident after load, working = extra peak per job).
DOCLING = DoclingEngine("ocrmac", EngineResources(model_memory_mb=1000, estimated_working_memory_mb=1800,
                                                   max_parallel_jobs=3))  # one converter, shared by threads
VISION = VisionEngine()
PADDLE = VlmEngine("paddleocr-vl-4bit", "paddle", "layout",
                   EngineResources(model_memory_mb=1400, estimated_working_memory_mb=1500, max_parallel_jobs=1))
# Whole page downscaled to 1024 px: 1.9 s/page instead of 13.4 at full size, and the lowest CER
# of all engines on degraded scans — the right last resort (REPORT.md).
GLM = VlmEngine("glm-ocr-4bit", "glm", "page",
                EngineResources(model_memory_mb=1600, estimated_working_memory_mb=1400, max_parallel_jobs=1),
                max_side=1024)
ENGINES = {router.DOCLING: DOCLING, router.VISION: VISION, router.PADDLE: PADDLE, router.GLM: GLM}
# If an engine isn't installed (models not downloaded, helper not built) use the next best one.
FALLBACK = {router.VISION: router.DOCLING, router.PADDLE: router.GLM, router.GLM: router.PADDLE}

VISION_CHUNK = 8   # pages per Vision helper call; chunks run in parallel up to VISION's max_parallel
VLM_GROUP = 4      # pages per VLM lease (the model stays loaded across them)
BLANK = "pagina vuota"


def usable(key: str) -> str | None:
    if ENGINES[key].available()[0]:
        return key
    alt = FALLBACK.get(key)
    if alt and ENGINES[alt].available()[0]:
        return alt
    return router.DOCLING if key == router.VISION else None


class Cancelled(Exception):
    pass


@dataclass
class Context:
    out_dir: Path                       # final document folder (published atomically at the end)
    work_dir: Path                      # internal: checkpoints, assets in progress, report
    name: str                           # original file name; the Markdown is <stem>.md
    mode: str = router.AUTO
    progress: Callable[[dict], None] = lambda _: None
    cancelled: Callable[[], bool] = lambda: False
    engines_used: set[str] = field(default_factory=set)
    notes: list[str] = field(default_factory=list)  # internal report only
    resumed_pages: int = 0                          # pages taken from checkpoints after a restart

    def check(self) -> None:
        if self.cancelled():
            raise Cancelled()

    @property
    def stem(self) -> str:
        return safe_name(Path(self.name).stem)


def safe_name(name: str) -> str:
    """File-system safe, still readable: keeps letters (any script), digits, spaces and . _ - ( )."""
    name = "".join(c if c.isalnum() or c in " ._-()" else "_" for c in name).strip(" .")
    return name or "documento"


def _runs(pages: list[int], size: int = CHUNK_PAGES) -> list[list[int]]:
    """Group sorted page numbers into contiguous runs of at most `size` pages."""
    runs: list[list[int]] = []
    for p in sorted(pages):
        if runs and p == runs[-1][-1] + 1 and len(runs[-1]) < size:
            runs[-1].append(p)
        else:
            runs.append([p])
    return runs


def image_to_pdf(src: Path, dst: Path) -> Path:
    """Wrap an image as a PDF sized like A4 width, so every engine sees a normal page."""
    with Image.open(src) as im:
        frames = []
        try:
            for i in range(getattr(im, "n_frames", 1)):  # multi-page TIFF
                im.seek(i)
                frames.append(im.convert("RGB"))
        except EOFError:
            pass
        dpi = frames[0].width / 8.27
        frames[0].save(dst, "PDF", save_all=True, append_images=frames[1:], resolution=dpi)
    return dst


def _ink(path: Path) -> float:
    """Share of dark pixels: tells a blank scanned page from one worth OCRing."""
    with Image.open(path) as im:
        return float((np.asarray(im.convert("L").resize((200, 280))) < 128).mean())


def _score(out: PageOutput, lang: str | None) -> None:
    out.quality = text_quality(out.markdown)
    ratio, words = misspelled_ratio(out.markdown, lang)
    out.data["misspelled"] = ratio
    out.data["words"] = words


def _record(out: PageOutput, tried: set[str]) -> dict:
    return {"page": out.page_no, "markdown": out.markdown, "engine": out.engine, "quality": out.quality,
            "data": out.data, "tried": sorted(tried)}


def _from_record(rec: dict) -> tuple[PageOutput, set[str]]:
    return (PageOutput(rec["page"], rec["markdown"], rec["engine"], rec.get("quality"), rec.get("data") or {}),
            set(rec.get("tried") or []))


# --------------------------------------------------------------------------- PDF / images
def convert_pdf(pdf: Path, info: FileInfo, ctx: Context, tmp: Path, assets: Path, cp: Checkpoint) -> list[PageOutput]:
    kinds = {p.number: p.kind for p in info.pages}
    avail = {k: usable(k) for k in ENGINES}
    total = len(kinds)
    results: dict[int, PageOutput] = {}
    tried: dict[int, set[str]] = {}
    for n, rec in cp.load().items():  # pages finished before an interruption
        if n in kinds:
            results[n], tried[n] = _from_record(rec)
    if results:
        ctx.notes.append(f"ripresa: {len(results)} pagine già completate caricate dal checkpoint")
    ctx.resumed_pages = len(results)
    plan = {n: avail[router.first_engine(k, ctx.mode)] or router.DOCLING
            for n, k in kinds.items() if n not in results}
    rendered: dict[int, Path] = {}
    done: set[int] = set(results)
    done_lock = threading.Lock()  # chunk threads add to `done` while others report progress

    def mark_done(pages: list[int]) -> None:
        with done_lock:
            done.update(pages)

    def report(stage: str) -> None:
        with done_lock:
            n = len(done)
        ctx.progress({"pages_done": n, "pages_total": total, "stage": stage})

    def image(n: int) -> Path:
        if n not in rendered:
            rendered[n] = render_pages(pdf, [n], tmp / "pages", config.OCR_DPI)[0]
        return rendered[n]

    def checkpoint(outs: list[PageOutput], key: str) -> None:
        for out in outs:
            cp.save(_record(out, tried.get(out.page_no, set()) | {key}))

    def run_docling(pages: list[int]) -> dict[int, PageOutput]:
        got = {}
        for run in _runs(pages):
            ctx.check()
            report(f"pagine {run[0]}-{run[-1]}")
            res = DOCLING.convert(pdf, run, cancelled=ctx.cancelled)
            outs = {o.page_no: o for o in page_outputs(res, DOCLING.name, run, assets)}
            del res  # don't keep chunk documents (and their pictures) alive on long PDFs
            for n in run:
                got[n] = outs.get(n) or PageOutput(n, "", DOCLING.name)
                _score(got[n], info.language)
            checkpoint([got[n] for n in run], router.DOCLING)
            mark_done(run)
            report(f"pagine {run[0]}-{run[-1]} completate")
        return got

    def run_ocr(key: str, pages: list[int]) -> dict[int, PageOutput]:
        """Image engines; blank pages are skipped, chunks run concurrently under memory leases."""
        got, todo = {}, []
        for n in pages:
            ctx.check()
            if _ink(image(n)) < 0.002:
                got[n] = PageOutput(n, "", BLANK)
                _score(got[n], info.language)
                checkpoint([got[n]], key)
            else:
                todo.append(n)
        size = VISION_CHUNK if key == router.VISION else VLM_GROUP
        chunks = [todo[i:i + size] for i in range(0, len(todo), size)]

        def one(chunk: list[int]) -> list[PageOutput]:
            ctx.check()
            report(f"pagine {chunk[0]}-{chunk[-1]}")
            imgs = [image(n) for n in chunk]
            outs = (VISION.convert_images(imgs, chunk) if key == router.VISION
                    else ENGINES[key].convert_images(imgs, chunk, assets))
            for out in outs:
                _score(out, info.language)
            checkpoint(outs, key)
            mark_done(chunk)
            report(f"pagine {chunk[0]}-{chunk[-1]} completate")
            return outs

        workers = ENGINES[key].resources.max_parallel_jobs
        with ThreadPoolExecutor(max_workers=max(1, min(workers, len(chunks)))) as pool:
            for outs in pool.map(one, chunks):
                for out in outs:
                    got[out.page_no] = out
        return got

    def run(key: str, pages: list[int]) -> dict[int, PageOutput]:
        if not pages:
            return {}
        ctx.engines_used.add(ENGINES[key].name)
        return run_docling(pages) if key == router.DOCLING else run_ocr(key, pages)

    # 1. First pass, engine by engine (cheapest first so results/progress arrive early).
    for key in (router.DOCLING, router.VISION, router.PADDLE):
        got = run(key, [n for n, e in plan.items() if e == key])
        results.update(got)
        for n in got:
            tried.setdefault(n, set()).add(key)
        report("verifica qualità")

    # 2. Escalation waves: Vision -> PaddleOCR-VL -> GLM-OCR, each page at most once per engine.
    #    Each page keeps exactly one result: the better of old and new (never both).
    for _ in range(3):
        wanted: dict[str, list[int]] = {}
        reasons: dict[int, str] = {}
        for n, out in results.items():
            if out.engine == BLANK:
                continue
            tables = sum(1 for b in out.data.get("blocks", []) if b.get("kind") == "table")
            esc = router.escalation(kinds[n], ctx.mode, _key(out.engine), out.markdown, out.quality,
                                    out.data.get("misspelled", 0.0), out.data.get("words", 0), tables)
            target = avail.get(esc[0]) if esc else None
            if target and target not in tried[n]:
                wanted.setdefault(target, []).append(n)
                reasons[n] = esc[1]
        if not wanted:
            break
        for key, pages in wanted.items():
            report(f"{len(pages)} pagine da rileggere con un motore più accurato")
            for n, new in run(key, pages).items():
                tried[n].add(key)
                old = results[n]
                if key == router.GLM and old.data.get("blocks"):
                    # GLM page mode is plain text: keep the layout model's table structure and headings.
                    new.markdown, notes = restore_tables(new.markdown, old.data["blocks"])
                    new.markdown = transfer_headings(new.markdown, old.data["blocks"])
                    new.data["table_notes"] = notes
                    new.data["blocks"] = old.data["blocks"]
                new.data["escalated_from"] = {"engine": old.engine, "reason": reasons[n],
                                              "misspelled": old.data.get("misspelled")}
                if router.better(_metrics(old), _metrics(new), reasons[n]):
                    results[n] = new
                else:
                    old.data["escalation_rejected"] = {"engine": new.engine, "reason": reasons[n]}
                cp.save(_record(results[n], tried[n]))
    report("assemblaggio del documento")
    return [results[n] for n in sorted(results)]


def _metrics(out: PageOutput) -> dict:
    return {"quality": out.quality or 0.0, "misspelled": out.data.get("misspelled", 0.0),
            "words": out.data.get("words", 0)}


def _key(engine_name: str) -> str:
    for key, eng in ENGINES.items():
        if eng.name == engine_name:
            return key
    return router.GLM  # blank pages: nothing to escalate


# --------------------------------------------------------------------------- charts
_GENERIC_ALT = {"", "image", "chart", "figure", "picture", "immagine", "grafico"}


def enrich_images(pages: list[PageOutput], base: Path, ctx: Context) -> set[str]:
    """Read the text printed in every image of the document; for charts with readable values add
    a verified data table after the image. Returns the image refs recognised as charts."""
    refs = [ref for p in pages for ref in markdown.image_refs(p.markdown) if (base / ref).is_file()]
    refs = list(dict.fromkeys(refs))
    if not refs or not VISION.available()[0]:
        return set()
    ctx.progress({"stage": "lettura di immagini e grafici"})
    texts: dict[str, list[str]] = {}
    for i in range(0, len(refs), 16):
        chunk = refs[i:i + 16]
        for ref, res in zip(chunk, VISION.recognize([base / r for r in chunk])):
            texts[ref] = [line["text"] for line in res.get("lines", [])]
    candidates = [r for r in refs if charts.numeric_tokens(texts[r]) >= 2]
    tables: dict[str, list[list[str]]] = {}
    chart_refs: set[str] = set()
    if candidates and PADDLE.available()[0]:
        imgs = []
        for r in candidates:
            with Image.open(base / r) as im:
                imgs.append(im.convert("RGB"))
        for ref, raw in zip(candidates, PADDLE.chart_tables(imgs)):
            if raw is None:
                continue
            chart_refs.add(ref)
            rows, reason = charts.verify_chart_table(raw, texts[ref])
            ctx.notes.append(f"grafico {Path(ref).name}: dati {reason}")
            if rows:
                tables[ref] = rows
        ctx.engines_used.add(PADDLE.name)

    def sub(m: re.Match) -> str:
        alt, ref = m.group(1), m.group(2)
        if alt.strip().lower() in _GENERIC_ALT and texts.get(ref):
            alt = charts.alt_text(texts[ref]).replace("]", ")").replace("[", "(")
        out = f"![{alt}]({ref})"
        if ref in tables:
            out += "\n\n<!-- dati letti dal grafico: ogni valore è stampato nell'immagine -->\n\n" \
                   + markdown.render_table(tables[ref])
        return out
    for p in pages:
        p.markdown = re.sub(r"!\[([^\]]*)\]\(([^)\s]+)\)", sub, p.markdown)
    return chart_refs


# --------------------------------------------------------------------------- other formats
def _demote(md: str, by: int) -> str:
    return re.sub(r"^(#{1,6})(?=\s)", lambda m: "#" * min(6, len(m.group(1)) + by), md, flags=re.M)


def convert_other(src: Path, info: FileInfo, ctx: Context, tmp: Path, assets: Path) -> list[PageOutput]:
    path = src
    if info.family == "text":
        return [PageOutput(0, src.read_text(encoding="utf-8", errors="replace"), "testo")]
    if info.family == "legacy":
        path = _legacy_to_modern(src, tmp)
    ctx.progress({"stage": "conversione"})
    res = DOCLING.convert(path, cancelled=ctx.cancelled)
    ctx.engines_used.add(DOCLING.name)
    outs = page_outputs(res, DOCLING.name, None, assets)
    ext = path.suffix.lower()
    if ext == ".pptx":
        notes = pptx_notes(path)
        for out in outs:
            note = notes.get(out.page_no)
            out.markdown = f"## Slide {out.page_no}\n\n{_demote(out.markdown, 2)}" + (
                f"\n\n> **Note del relatore:** {note}" if note else "")
    elif ext in (".xlsx", ".xlsm"):
        sheets = xlsx_details(path)
        for out, sheet in zip(outs, sheets):
            out.markdown = f"## Foglio: {sheet['name']}\n\n{_demote(out.markdown, 2)}"
            if sheet["formulas"]:
                out.markdown += "\n\n**Formule:**\n\n" + "\n".join(
                    f"- `{f['cell']}`: `{f['formula']}`" + (f" = {f['value']}" if f["value"] is not None else "")
                    for f in sheet["formulas"])
            out.data["sheet"] = sheet
    return outs


def _legacy_to_modern(src: Path, tmp: Path) -> Path:
    ext = src.suffix.lower()
    office = soffice()
    target = {".doc": "docx", ".rtf": "docx", ".odt": "docx", ".ppt": "pptx", ".odp": "pptx",
              ".xls": "xlsx", ".ods": "xlsx"}[ext]
    out = tmp / f"{src.stem}.{target}"
    if office:
        cmd = [office, "--headless", "--convert-to", target, "--outdir", str(tmp), str(src)]
    elif target == "docx":
        cmd = ["/usr/bin/textutil", "-convert", "docx", "-output", str(out), str(src)]
    else:
        raise RuntimeError(f"Il formato {ext} richiede LibreOffice")
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    if proc.returncode != 0 or not out.exists():
        raise RuntimeError(f"Conversione {ext} → {target} fallita: {proc.stderr.strip()[:300]}")
    return out


# --------------------------------------------------------------------------- metadata
_PLACEHOLDER = {"anonymous", "(anonymous)", "unknown", "untitled", "python-docx", "python-pptx", "openpyxl",
                "unspecified", "(unspecified)"}


def source_metadata(src: Path, info: FileInfo) -> dict[str, str]:
    """Title/author/date as recorded in the source file itself; generator defaults are ignored."""
    meta: dict[str, str] = {}
    try:
        if info.family == "pdf":
            import pypdfium2 as pdfium
            from docling.utils.locks import pypdfium2_lock
            with pypdfium2_lock:
                doc = pdfium.PdfDocument(src)
                try:
                    md = doc.get_metadata_dict()
                finally:
                    doc.close()
            meta = {"title": md.get("Title", ""), "author": md.get("Author", ""), "subject": md.get("Subject", "")}
            if m := re.match(r"D:(\d{4})(\d{2})(\d{2})", md.get("CreationDate", "")):
                meta["date"] = "-".join(m.groups())
        elif src.suffix.lower() in (".docx", ".pptx"):
            import docx
            import pptx
            props = (docx.Document(src) if src.suffix.lower() == ".docx" else pptx.Presentation(src)).core_properties
            meta = {"title": props.title or "", "author": props.author or "", "subject": props.subject or "",
                    "date": props.created.date().isoformat() if props.created else ""}
        elif src.suffix.lower() in (".xlsx", ".xlsm"):
            import openpyxl
            props = openpyxl.load_workbook(src, read_only=True).properties
            meta = {"title": props.title or "", "author": props.creator or "",
                    "date": props.created.date().isoformat() if props.created else ""}
    except Exception as exc:  # metadata is optional: never fail a conversion for it
        log.debug("metadata unreadable for %s: %s", src, exc)
    stems = {src.stem.lower(), Path(info.path).stem.lower()}
    return {k: v.strip() for k, v in meta.items() if v and v.strip() and v.strip().lower() not in _PLACEHOLDER
            and not (k == "title" and v.strip().lower() in stems)}  # a "title" that is just the file name


# --------------------------------------------------------------------------- publish
def publish(md: str, pages: list[PageOutput], info: FileInfo, ctx: Context, chart_refs: set[str],
            seconds: float) -> Path:
    """Write <stem>.md + assets/ + metadata.json into a staging dir, then swap it into place, so
    a reader never sees a half-written result. Only images the Markdown links are published."""
    stage = ctx.work_dir / "final"
    shutil.rmtree(stage, ignore_errors=True)
    (stage / "assets").mkdir(parents=True)
    mapping: dict[str, str | None] = {}
    counters = {"grafico": 0, "immagine": 0}
    for ref in dict.fromkeys(markdown.image_refs(md)):
        src = ctx.work_dir / ref
        if not src.is_file():
            mapping[ref] = None
            continue
        kind = "grafico" if ref in chart_refs else "immagine"
        counters[kind] += 1
        name = f"{kind}_{counters[kind]:03d}{src.suffix.lower()}"
        shutil.copy2(src, stage / "assets" / name)
        mapping[ref] = f"assets/{name}"
    md = markdown.rewrite_images(md, mapping)
    if not any((stage / "assets").iterdir()):
        (stage / "assets").rmdir()
    atomic_write(stage / f"{ctx.stem}.md", md)
    metadata = {"source": ctx.name, "format": info.ext.lstrip(".") or info.mime, "size_bytes": info.size,
                "pages": info.page_count or len(pages), "mode": ctx.mode, "language": info.language,
                "converted_at": datetime.now().astimezone().isoformat(timespec="seconds"),
                "duration_s": round(seconds, 1)}
    sheets = [p.data["sheet"] for p in pages if "sheet" in p.data]
    if sheets:
        metadata["sheets"] = [{"name": s["name"], "range": s["dimensions"], "merged_cells": s["merged"]}
                              for s in sheets]
    atomic_write(stage / "metadata.json", json.dumps(metadata, ensure_ascii=False, indent=2))
    if ctx.out_dir.exists():
        shutil.rmtree(ctx.out_dir)
    ctx.out_dir.parent.mkdir(parents=True, exist_ok=True)
    os.replace(stage, ctx.out_dir)
    return ctx.out_dir / f"{ctx.stem}.md"


# --------------------------------------------------------------------------- entry point
def convert_file(src: Path, ctx: Context) -> dict:
    """Convert `src` into ctx.out_dir (<stem>.md, assets/, metadata.json). Returns a summary with
    the internal per-page report (also saved as work_dir/report.json, never published)."""
    t0 = time.time()
    info = analyze(src)
    ctx.progress({"stage": "analisi", "pages_total": info.page_count})
    if info.error:
        raise ValueError(info.error)
    ctx.work_dir.mkdir(parents=True, exist_ok=True)
    if shutil.disk_usage(ctx.work_dir).free < config.MIN_FREE_DISK:
        raise OSError("Spazio su disco insufficiente (meno di 2 GB liberi)")
    assets = ctx.work_dir / "assets"
    config.TEMP.mkdir(parents=True, exist_ok=True)
    with Checkpoint(ctx.work_dir, src, ctx.mode) as cp, tempfile.TemporaryDirectory(dir=config.TEMP) as td:
        tmp = Path(td)
        if info.family in ("pdf", "image"):
            pdf = src if info.family == "pdf" else image_to_pdf(src, tmp / "image.pdf")
            if info.family == "image":
                info = _reanalyze_image(pdf, info)
            pages = convert_pdf(pdf, info, ctx, tmp, assets, cp)
        else:
            pages = convert_other(src, info, ctx, tmp, assets)
        ctx.check()
        chart_refs = enrich_images(pages, ctx.work_dir, ctx) if ctx.mode != router.FAST else set()
        md = markdown.assemble([p.markdown for p in pages], source_metadata(src, info))
        md_path = publish(md, pages, info, ctx, chart_refs, time.time() - t0)
        cp.clear()
    report = {
        "source": ctx.name, "mode": ctx.mode, "seconds": round(time.time() - t0, 2),
        "engines": sorted(ctx.engines_used), "analysis": info.as_dict(), "notes": ctx.notes,
        "resumed_pages": ctx.resumed_pages,
        "escalated_pages": [p.page_no for p in pages if "escalated_from" in p.data],
        "pages": [{"page": p.page_no, "engine": p.engine, "quality": p.quality,
                   **{k: v for k, v in p.data.items() if k != "blocks"}} for p in pages],
    }
    atomic_write(ctx.work_dir / "report.json", json.dumps(report, ensure_ascii=False, indent=1, default=str))
    shutil.rmtree(assets, ignore_errors=True)  # published copies live in out_dir/assets
    return {"pages": len(pages), "page_kinds": info.counts(), "language": info.language,
            "seconds": report["seconds"], "engines": report["engines"], "resumed_pages": ctx.resumed_pages,
            "escalated_pages": report["escalated_pages"], "markdown": str(md_path), "report": report}


def _reanalyze_image(pdf: Path, info: FileInfo) -> FileInfo:
    from app.analyzer import analyze_pdf
    info.pages = []  # replaced by the wrapped PDF's pages (several for a multi-page TIFF)
    analyze_pdf(pdf, info)
    for p in info.pages:
        p.kind = "scanned"
    return info
